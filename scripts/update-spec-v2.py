# -*- coding: utf-8 -*-
"""Actualiza la spec: foto de identidad como evidencia + grabación de video.
Parte de la versión ACTUAL (con tus ediciones) y guarda una v2, sin tocar tu archivo."""
import shutil
from docx import Document

SRC = r"C:\Users\tomas\Documents\ProctoTeam-UGM\Especificación Funcional - ProctoTeam UGM (con contenido).docx"
DST = r"C:\Users\tomas\Documents\ProctoTeam-UGM\Especificación Funcional - ProctoTeam UGM (con contenido) v2.docx"

shutil.copyfile(SRC, DST)
doc = Document(DST)


def replace_para(needle, new_text):
    for p in doc.paragraphs:
        if needle in p.text:
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            p.add_run(new_text)
            print("OK:", needle[:42])
            return True
    print("** NO MATCH:", needle[:42])
    return False


# 1.3 — bullet de verificación de identidad (foto guardada como evidencia)
replace_para(
    "Verificación de identidad al ingreso",
    "• Verificación de identidad al ingreso: el estudiante toma una fotografía de sí mismo sosteniendo su documento de identidad (cédula/carnet). El sistema lee el nombre del documento mediante IA (OCR) y el estudiante lo confirma, validándolo contra el alumno matriculado. Esa fotografía se almacena como evidencia para su revisión posterior.",
)

# 1.3 — consideración sobre la evidencia (ahora SÍ graba video)
replace_para(
    "Consideración sobre la evidencia",
    "Consideración sobre la evidencia: ProctoTeam graba la sesión de examen y la almacena como evidencia, junto con la fotografía de verificación de identidad (el estudiante con su documento) y las capturas de las alertas graves. Esto permite revisar cada evaluación a posteriori. El volumen asociado a la grabación se dimensiona y gestiona mediante políticas de retención (ver Diseño Técnico).",
)

# 2 — Almacenamiento de evidencia (3 tipos: identidad, alertas, grabación)
replace_para(
    "Las evidencias de las alertas graves se almacenan",
    "El sistema almacena en la nube (Supabase Storage) tres tipos de evidencia: (1) la fotografía de verificación de identidad (el estudiante con su documento), capturada al ingreso; (2) las capturas (imágenes JPEG) asociadas a las alertas graves; y (3) la grabación de video de la sesión de examen. Cada archivo se nombra de forma trazable (identificador de participación, tipo y marca de tiempo) y su referencia (URL) se registra en la base de datos, vinculando la evidencia con el estudiante y, cuando corresponde, con el incidente.",
)

# 2 — Tiempo real y video (añade que la sesión se graba)
replace_para(
    "Esta arquitectura P2P descarga",
    "Esta arquitectura P2P descarga el tráfico de video del servidor, reduciendo costos y latencia, a costa de concentrar en el navegador del instructor la recepción de múltiples flujos simultáneos. Adicionalmente, la sesión se graba y se almacena como evidencia, lo que permite su revisión posterior aunque el instructor no haya observado ese momento en vivo.",
)

# 2 — Capa de IA en el servidor (OCR, no comparación facial)
replace_para(
    "En el servidor: flujos de IA",
    "• En el servidor: flujos de IA basados en modelos de lenguaje y visión (sobre OpenAI GPT, mediante el AI SDK de Vercel) que realizan la lectura del documento de identidad (extracción del nombre por OCR/visión), la evaluación de severidad de las alertas y los asistentes de ayuda. Las salidas se entregan en formato estructurado validado con esquemas (Zod).",
)

# 2 — Nota de cumplimiento (OCR evita reconocimiento facial)
replace_para(
    "Nota de cumplimiento",
    "Nota de cumplimiento: la verificación de identidad se basa en la lectura del nombre del documento (OCR), evitando el reconocimiento facial automatizado. La fotografía del estudiante con su documento se conserva como evidencia para eventual revisión humana. Si en el futuro se incorpora comparación facial automática, se evaluará un servicio biométrico especializado conforme a las políticas aplicables.",
)

# 2 — Flujo del estudiante (verificación = foto + confirmar nombre)
replace_para(
    "Realiza la verificación",
    "• Realiza la verificación de identidad: toma una fotografía sosteniendo su documento (que se guarda como evidencia) y confirma su nombre leído del documento.",
)

# 2 — Seguridad y privacidad (sesión grabada; evidencia incluye grabación)
replace_para(
    "Toda la comunicación viaja cifrada",
    "Toda la comunicación viaja cifrada: HTTPS para la aplicación y DTLS-SRTP para los flujos de WebRTC. El acceso a la información está controlado por rol. El estudiante es informado de que la sesión es monitoreada y grabada (audio y video). La evidencia almacenada (fotografía de identidad, capturas de alertas y grabación de la sesión) se gestiona bajo control de acceso por rol y políticas de retención. Se recomienda definir políticas explícitas de consentimiento informado y de retención/eliminación de evidencias acordes a la normativa de protección de datos aplicable.",
)

# 2 — Estimación de almacenamiento (ahora dominada por la grabación)
replace_para(
    "Dado que ProctoTeam almacena",
    "Dado que ProctoTeam graba la sesión de examen, el principal componente de almacenamiento es el video. A modo de referencia, una evaluación de 30 estudiantes durante 60 minutos requiere del orden de 13 a 30 GB: una grabación de webcam a calidad estándar ronda los 13 GB, y al incluir la pantalla compartida se acerca a 25-30 GB. A esto se suma un volumen menor por las fotografías de identidad y las capturas de alertas (algunos MB por sesión). Estos volúmenes deben considerarse en los costos de almacenamiento y en las políticas de retención.",
)

# 2 — Limitaciones (reemplaza "no existe grabación" por una real)
replace_para(
    "No existe grabación de video continua",
    "• La calidad y la continuidad de la grabación dependen del hardware y del ancho de banda del estudiante; cortes de conexión pueden producir vacíos en el video.",
)

# 2 — Trabajo futuro (quita "grabación opcional", ya es parte de la solución)
replace_para(
    "Trabajo futuro: grabación opcional",
    "• Trabajo futuro: proveedor biométrico especializado para verificación facial, reportes descargables e informes automatizados, controles de retención configurables y analítica de integridad.",
)

doc.save(DST)
print("GUARDADO:", DST)
