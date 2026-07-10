# -*- coding: utf-8 -*-
"""Rellena la Especificación Funcional de ProctoTeam en una COPIA del .docx original."""
import shutil
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

SRC = r"C:\Users\tomas\Documents\ProctoTeam-UGM\Especificación Funcional - ProcoTeam UGM.docx"
DST = r"C:\Users\tomas\Documents\ProctoTeam-UGM\Especificación Funcional - ProctoTeam UGM (con contenido).docx"

shutil.copyfile(SRC, DST)
doc = Document(DST)
styles = set(s.name for s in doc.styles)
SUB = 'Heading 3' if 'Heading 3' in styles else 'Normal'


def insert_after(anchor, text, style):
    new_p = OxmlElement('w:p')
    anchor._p.addnext(new_p)
    para = Paragraph(new_p, anchor._parent)
    st = style
    if style == 'List Bullet' and 'List Bullet' not in styles:
        text = '• ' + text
        st = 'Normal'
    if st not in styles:
        st = 'Normal'
    para.style = doc.styles[st]
    if text:
        para.add_run(text)
    return para


def fill(heading_text, blocks):
    target = None
    for p in doc.paragraphs:
        if p.text.strip().lower() == heading_text.strip().lower():
            target = p
            break
    if target is None:
        key = heading_text.strip().lower()[:18]
        for p in doc.paragraphs:
            if p.text.strip().lower().startswith(key):
                target = p
                break
    if target is None:
        print("NO ENCONTRADO:", heading_text)
        return
    anchor = target
    for style, text in blocks:
        anchor = insert_after(anchor, text, style)
    print("OK:", heading_text, "->", len(blocks), "parrafos")


P = 'Normal'
B = 'List Bullet'

S11 = [
    (P, "La Universidad Gabriela Mistral (UGM) requiere implementar un servicio de proctoring (supervisión remota de evaluaciones en línea) que permita resguardar la integridad académica y la calidad educativa de los exámenes rendidos bajo modalidad online. El crecimiento sostenido de la evaluación a distancia, impulsado por la transformación digital de la institución, exige garantizar que las calificaciones obtenidas de forma remota tengan la misma validez y confiabilidad que las presenciales."),
    (P, "Desde la perspectiva del usuario, el requerimiento se descompone en las siguientes necesidades:"),
    (B, "Verificar la identidad del estudiante antes de iniciar la evaluación, asegurando que quien rinde el examen es efectivamente la persona matriculada."),
    (B, "Supervisar en tiempo real el desarrollo del examen, permitiendo al docente observar a los estudiantes durante toda la sesión."),
    (B, "Detectar de forma automática conductas irregulares (ausencia frente a la cámara, presencia de terceros, uso de dispositivos no autorizados, cambio de ventana, etc.)."),
    (B, "Generar evidencia objetiva y trazable de los incidentes, que respalde eventuales procesos académicos o disciplinarios."),
    (B, "Entregar a los docentes herramientas de monitoreo y gestión simples, y a la institución indicadores y trazabilidad a nivel transversal."),
    (P, "En síntesis, la UGM necesita una solución que combine verificación de identidad, vigilancia en vivo, detección asistida por inteligencia artificial y registro de evidencia, de modo de proteger la credibilidad de sus procesos evaluativos y, con ello, el valor de los títulos que otorga."),
]

S12 = [
    (P, "El problema central que aborda este proyecto es la copia y la suplantación de identidad en las evaluaciones de modalidad online. A diferencia de un examen presencial, donde un supervisor controla físicamente la sala, en la modalidad remota el estudiante rinde desde un entorno no controlado, lo que abre múltiples vectores de fraude académico:"),
    (B, "Suplantación de identidad: que una persona distinta al estudiante matriculado rinda la evaluación en su lugar."),
    (B, "Copia y uso de material no autorizado: apuntes, sitios web, aplicaciones o un segundo dispositivo fuera del campo de la cámara."),
    (B, "Ayuda externa: presencia de terceros en la sala o asistencia remota durante el examen."),
    (B, "Abandono del entorno de evaluación: cambiar de pestaña o ventana para buscar respuestas en línea."),
    (P, "El impacto de estas conductas es significativo: erosionan la validez de las calificaciones, generan inequidad entre los estudiantes que actúan con honestidad y quienes no, y exponen a la institución a riesgos reputacionales y de acreditación, en la medida en que la confianza en los resultados de aprendizaje se ve comprometida."),
    (P, "Los enfoques tradicionales resultan insuficientes: confiar únicamente en la honestidad del estudiante no ofrece garantías, mientras que exigir exámenes presenciales para toda evaluación es costoso, logísticamente complejo y contrario a la estrategia de digitalización de la UGM. Se requiere, por tanto, un mecanismo tecnológico que disuada y detecte el fraude sin renunciar a las ventajas de la modalidad online."),
]

S13 = [
    (P, "La solución propuesta es ProctoTeam, una plataforma web de proctoring desarrollada a la medida de la UGM, que integra autenticación institucional, verificación biométrica de identidad, monitoreo en vivo y detección de irregularidades asistida por inteligencia artificial."),
    (P, "Alcance: la solución está concebida para toda la UGM, abarcando la totalidad de carreras y facultades, y a los perfiles de usuario principales (estudiantes e instructores), además de un rol de administración institucional (super-admin). ProctoTeam se integra con la plataforma de evaluación Blackboard que ya utiliza la universidad: el examen propiamente tal se rinde en Blackboard, mientras ProctoTeam provee la capa de supervisión."),
    (P, "Principales funcionalidades:"),
    (B, "Autenticación institucional mediante inicio de sesión único (SSO) con Microsoft Azure AD / Entra ID, usando las cuentas UGM. El sistema gestiona tres roles: estudiante, instructor y administrador (super-admin)."),
    (B, "Verificación biométrica de identidad al ingreso: el estudiante captura una foto de su carnet o cédula y una foto en vivo de su rostro, y un modelo de inteligencia artificial compara ambas imágenes para confirmar que corresponden a la misma persona."),
    (B, "Monitoreo en vivo por WebRTC: cada estudiante transmite su cámara, su pantalla compartida y su micrófono al instructor en tiempo real. El instructor visualiza simultáneamente a todos los estudiantes de la sesión en una grilla."),
    (B, "Detección con IA en el navegador del estudiante: mediante TensorFlow.js y MediaPipe se detectan localmente situaciones como ausencia del rostro, presencia de múltiples personas, objetos sospechosos (por ejemplo, teléfonos), cambio de pestaña o pérdida del modo pantalla completa."),
    (B, "Sistema de alertas con niveles de severidad. Ante alertas graves, el sistema captura automáticamente una evidencia (imagen compuesta de la pantalla y la webcam) que se almacena en la nube y se registra en la base de datos, asociada al estudiante y al tipo de incidente."),
    (B, "Asistentes de ayuda con IA, diferenciados para estudiante e instructor, que responden dudas sobre el uso de la plataforma."),
    (B, "Paneles de gestión: dashboard del instructor (monitoreo en vivo y alertas), histórico de sesiones, y panel de super-admin para la gestión de instructores y métricas institucionales."),
    (B, "Integración con Blackboard: la evaluación se abre y rinde en la plataforma institucional, manteniendo ProctoTeam como capa de vigilancia."),
    (P, "Consideración sobre la evidencia: ProctoTeam conserva la transmisión en vivo de forma efímera (no almacena grabaciones de video continuas) y persiste imágenes (capturas) como evidencia de las alertas graves. Esto reduce drásticamente los costos de almacenamiento y los riesgos de privacidad, manteniendo a la vez un respaldo objetivo de los incidentes. Una eventual grabación de video completa se contempla como evolución futura (ver Diseño Técnico)."),
]

S2 = [
    (P, "Esta sección describe la arquitectura técnica de ProctoTeam, sus componentes, integraciones y decisiones de diseño. La plataforma se construye sobre tecnologías web modernas y servicios gestionados en la nube, priorizando la escalabilidad, la seguridad y un bajo costo operativo."),

    (SUB, "Arquitectura general"),
    (P, "ProctoTeam es una aplicación web construida con Next.js (React, App Router) y desplegada en la plataforma Vercel. La lógica de servidor se implementa mediante API Routes y Server Actions ejecutadas en funciones serverless, lo que evita administrar infraestructura propia y permite escalar bajo demanda. El frontend entrega la experiencia para los tres roles (estudiante, instructor y super-admin), mientras que la persistencia y los servicios en tiempo real se apoyan en Supabase (PostgreSQL, Storage y Realtime)."),

    (SUB, "Autenticación y autorización"),
    (P, "El inicio de sesión se realiza mediante SSO con Microsoft Azure AD / Entra ID, utilizando la librería MSAL y el protocolo OpenID Connect con flujo Authorization Code + PKCE. Tras autenticarse, el cliente obtiene un id_token que el backend valida criptográficamente: verifica la firma RS256 contra el JWKS del tenant institucional y comprueba el audience (identificador de la aplicación) y el tenant de procedencia."),
    (P, "La autorización se basa en roles almacenados en la base de datos. Por seguridad, todo usuario nuevo se crea con rol estudiante; los roles instructor y super-admin se asignan de forma controlada. Las rutas de la aplicación y los endpoints sensibles validan el rol antes de exponer información o permitir acciones."),

    (SUB, "Modelo de datos y persistencia"),
    (P, "La información se almacena en una base de datos PostgreSQL gestionada por Supabase. Las entidades principales son: usuarios (users), sesiones de examen (exam_sessions), participaciones de estudiantes (exam_participations) y alertas (alerts). El acceso se realiza mediante conexión directa a PostgreSQL con un mecanismo de respaldo a través de la API REST de Supabase, lo que aporta resiliencia ante distintas condiciones de red."),

    (SUB, "Almacenamiento de evidencia"),
    (P, "Las evidencias de las alertas graves se almacenan como imágenes JPEG en Supabase Storage, en un bucket dedicado. Cada archivo se nombra de forma trazable combinando el identificador de la participación, el tipo de alerta y una marca de tiempo. La referencia (URL) de cada evidencia se registra en la tabla de alertas, vinculando la imagen con el estudiante y el incidente correspondiente."),

    (SUB, "Tiempo real y video (WebRTC)"),
    (P, "El monitoreo en vivo se implementa con WebRTC en topología punto a punto (peer-to-peer) desde cada estudiante hacia el instructor. La señalización (intercambio de ofertas, respuestas y candidatos ICE) se realiza a través de los canales de Supabase Realtime (Broadcast), y se emplean servidores STUN para el establecimiento de la conexión. El estudiante transmite su cámara (resolución típica 640×480), su pantalla (capturada con la API de compartición de pantalla) y su micrófono."),
    (P, "Esta arquitectura P2P descarga el tráfico de video del servidor, reduciendo costos y latencia, a costa de concentrar en el navegador del instructor la recepción de múltiples flujos simultáneos."),

    (SUB, "Capa de inteligencia artificial"),
    (P, "ProctoTeam combina inteligencia artificial en el cliente y en el servidor:"),
    (B, "En el cliente (navegador del estudiante): se ejecutan modelos con TensorFlow.js (detección de objetos con COCO-SSD) y MediaPipe (malla facial) para detectar en tiempo real ausencia o múltiples rostros, objetos sospechosos y otras situaciones, sin necesidad de enviar el video al servidor, lo que protege la privacidad y reduce el ancho de banda."),
    (B, "En el servidor: flujos de IA basados en modelos de lenguaje y visión (actualmente sobre OpenAI GPT, mediante el AI SDK de Vercel) que realizan la verificación biométrica de identidad (comparación de las imágenes de carnet y rostro), la evaluación de severidad de las alertas y los asistentes de ayuda. Las salidas se entregan en formato estructurado validado con esquemas (Zod)."),
    (P, "Nota de cumplimiento: el uso de modelos de IA de propósito general para verificación biométrica está sujeto a las políticas del proveedor; para producción se evaluará el uso de un servicio biométrico especializado para ese flujo específico."),

    (SUB, "Flujo del estudiante"),
    (B, "Inicia sesión con su cuenta institucional UGM (SSO)."),
    (B, "Ingresa el código de acceso del examen para unirse a la sesión."),
    (B, "Realiza la verificación biométrica (foto del carnet + foto del rostro)."),
    (B, "Acepta los requisitos de monitoreo."),
    (B, "Se activa el monitoreo (cámara, pantalla compartida y detección por IA)."),
    (B, "Rinde el examen en Blackboard, abierto desde ProctoTeam."),
    (B, "Finaliza la sesión de monitoreo."),

    (SUB, "Flujo del instructor"),
    (B, "Inicia sesión en la plataforma."),
    (B, "Configura o crea el examen, lo que genera un código de acceso para compartir con los estudiantes."),
    (B, "Supervisa en el monitor en vivo (grilla de estudiantes y panel de alertas en tiempo real)."),
    (B, "Consulta el histórico de sesiones."),

    (SUB, "Seguridad y privacidad"),
    (P, "Toda la comunicación viaja cifrada: HTTPS para la aplicación y DTLS-SRTP para los flujos de WebRTC. El acceso a la información está controlado por rol. El estudiante es informado de que la sesión es monitoreada (audio y video). La evidencia se limita a capturas de incidentes, minimizando los datos personales almacenados. Se recomienda definir políticas explícitas de consentimiento y de retención/eliminación de evidencias acordes a la normativa aplicable."),

    (SUB, "Requisitos no funcionales"),
    (B, "Escalabilidad: el modelo P2P evita que el servidor procese video; el límite práctico está dado por la capacidad del navegador y la conexión del instructor para recibir N flujos simultáneos. Para cursos numerosos se recomienda acotar el tamaño de las salas o paginar la grilla de monitoreo."),
    (B, "Rendimiento: la detección por IA se ejecuta en el cliente para no sobrecargar el servidor."),
    (B, "Requisitos del cliente: navegador moderno (Chrome o Edge recomendados), cámara, micrófono y permisos de cámara/pantalla habilitados, y conexión estable."),

    (SUB, "Estimación de almacenamiento"),
    (P, "Dado que ProctoTeam almacena únicamente capturas de las alertas graves (no video continuo), el consumo por evaluación es del orden de pocos MB. A modo de referencia, si en el futuro se implementara grabación de video completa por estudiante, una evaluación de 30 estudiantes durante 60 minutos requeriría aproximadamente entre 13 y 30 GB (según la calidad y si se incluye la pantalla compartida), lo que debe considerarse en costos y políticas de retención antes de habilitar dicha funcionalidad."),

    (SUB, "Limitaciones conocidas y trabajo futuro"),
    (B, "No existe grabación de video continua; sólo se conservan capturas de las alertas y la transmisión en vivo es efímera."),
    (B, "El monitoreo en vivo requiere que el instructor esté presente observando; lo que no se observa en vivo no queda registrado en video."),
    (B, "La experiencia depende del hardware y la conexión del estudiante."),
    (B, "Trabajo futuro: grabación opcional de la sesión (por estudiante o de la vista del instructor), proveedor biométrico especializado, reportes descargables e informes automatizados, y reglas de retención de evidencia."),
]

fill("Descripción del requerimiento del usuario", S11)
fill("Descripción del problema", S12)
fill("Solución y alcance", S13)
fill("Diseño técnico", S2)

doc.save(DST)
print("GUARDADO:", DST)

d2 = Document(DST)
print("Parrafos totales tras rellenar:", len(d2.paragraphs))
